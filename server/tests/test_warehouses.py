import io

from openpyxl import Workbook


def _push(client, tx_id, qr, action, qty, ts, user="u1", warehouse=None):
    tx = {
        "tx_id": tx_id,
        "qr_id": qr,
        "user_id": user,
        "action_type": action,
        "timestamp": ts,
        "quantity": qty,
    }
    if warehouse is not None:
        tx["warehouse"] = warehouse
    return client.post(
        "/api/sync/push", json={"device_id": "d", "transactions": [tx]}
    ).json()


def _stock(client, qr):
    pulled = client.get("/api/sync/pull").json()
    item = next(i for i in pulled["items"] if i["qr_id"] == qr)
    shelves = {
        s["warehouse"]: s["quantity"] for s in pulled["item_stock"] if s["qr_id"] == qr
    }
    return item, shelves


def _setup(client):
    client.post("/api/admin/users", json=[{"user_id": "u1", "full_name": "א"}])
    for name in ("A", "B", "C"):
        assert client.post("/api/panel/warehouses", json={"name": name}).json()[
            "success"
        ]
    client.post(
        "/api/panel/items",
        json={
            "qr_id": "BAT",
            "category": "סוללות",
            "quantity": 1,
            "kind": "CONSUMABLE",
        },
    )
    r = client.post(
        "/api/panel/items/stock",
        json={"qr_id": "BAT", "shelves": {"A": 20, "B": 50, "C": 30}},
    )
    assert r.json()["success"], r.json()


def test_split_sets_stock_and_issues_come_off_one_warehouse(client):
    _setup(client)
    item, shelves = _stock(client, "BAT")
    assert item["quantity"] == 100 and shelves == {"A": 20, "B": 50, "C": 30}
    assert client.get("/api/sync/pull").json()["warehouses"] == ["A", "B", "C"]

    assert _push(client, "t1", "BAT", "ISSUE", 5, 1_000, warehouse="B")["accepted"] == [
        "t1"
    ]
    item, shelves = _stock(client, "BAT")
    assert item["quantity"] == 95 and shelves == {"A": 20, "B": 45, "C": 30}

    # More than that warehouse has is refused, even if the total would do.
    body = _push(client, "t2", "BAT", "ISSUE", 25, 2_000, warehouse="A")
    assert (
        body["accepted"] == [] and "אין מספיק במחסן A" in body["rejected"][0]["message"]
    )

    # Moving units between warehouses keeps the issued count.
    client.post(
        "/api/panel/items/stock",
        json={"qr_id": "BAT", "shelves": {"A": 40, "B": 25, "C": 30}},
    )
    item, shelves = _stock(client, "BAT")
    assert (item["quantity"], item["issued_qty"]) == (95, 5)
    assert shelves == {"A": 40, "B": 25, "C": 30}

    # A plain item edit does not change a split item's stock.
    client.post(
        "/api/panel/items",
        json={
            "qr_id": "BAT",
            "category": "סוללות",
            "quantity": 1,
            "kind": "CONSUMABLE",
        },
    )
    assert _stock(client, "BAT")[0]["quantity"] == 95


def test_loans_return_to_a_warehouse(client):
    _setup(client)
    client.post(
        "/api/panel/items", json={"qr_id": "K", "category": "קסדה", "quantity": 1}
    )
    client.post(
        "/api/panel/items/stock", json={"qr_id": "K", "shelves": {"A": 3, "B": 2}}
    )
    assert _stock(client, "K")[0]["quantity"] == 5
    assert _push(client, "b1", "K", "BORROW", 2, 1_000, warehouse="A")["accepted"] == [
        "b1"
    ]
    item, shelves = _stock(client, "K")
    assert (item["available_qty"], shelves) == (3, {"A": 1, "B": 2})
    assert _push(client, "r1", "K", "RETURN", 2, 2_000, warehouse="B")["accepted"] == [
        "r1"
    ]
    item, shelves = _stock(client, "K")
    assert (item["available_qty"], shelves) == (5, {"A": 1, "B": 4})
    # Older tablets send no warehouse: only the item total moves.
    assert _push(client, "b2", "K", "BORROW", 1, 3_000)["accepted"] == ["b2"]
    assert _stock(client, "K")[1] == {"A": 1, "B": 4}


def test_warehouse_rename_and_delete(client):
    _setup(client)
    _push(client, "t1", "BAT", "ISSUE", 5, 1_000, warehouse="A")
    r = client.post("/api/panel/warehouses", json={"name": "ראשי", "old_name": "A"})
    assert r.json()["success"]
    assert _stock(client, "BAT")[1] == {"ראשי": 15, "B": 50, "C": 30}
    history = client.get("/api/sync/pull").json()["history"]
    assert history[0]["warehouse"] == "ראשי"
    r = client.delete("/api/panel/warehouses", params={"name": "C"})
    assert r.status_code == 400 and "30" in r.json()["message"]
    client.post(
        "/api/panel/items/stock",
        json={"qr_id": "BAT", "shelves": {"ראשי": 15, "B": 80, "C": 0}},
    )
    assert client.delete("/api/panel/warehouses", params={"name": "C"}).json()[
        "success"
    ]
    data = client.get("/api/panel/data").json()
    assert {w["name"]: w["units"] for w in data["warehouses"]} == {"B": 80, "ראשי": 15}


def test_excel_import_splits_by_warehouse(client):
    wb = Workbook()
    ws = wb.active
    ws.append(["מספר סידורי", "סוג פריט", "סוג שימוש", "מחסן", "כמות"])
    ws.append(["AA", "סוללות AA", "ניצרך", "A", 20])
    ws.append(["AA", "סוללות AA", "ניצרך", "B", 50])
    out = io.BytesIO()
    wb.save(out)
    r = client.post(
        "/api/panel/import", content=out.getvalue(), headers={"X-Filename": "w.xlsx"}
    )
    assert r.json()["success"], r.json()
    item, shelves = _stock(client, "AA")
    assert (item["quantity"], shelves) == (70, {"A": 20, "B": 50})
    assert client.get("/api/panel/export/warehouses").status_code == 200


def test_shelves_and_stock_count(client):
    _setup(client)
    client.post(
        "/api/panel/items", json={"qr_id": "K", "category": "קסדה", "quantity": 1}
    )
    client.post(
        "/api/panel/items/stock",
        json={"qr_id": "K", "shelves": {"A": 3}, "shelf_numbers": {"A": "10"}},
    )
    client.post(
        "/api/panel/items/stock",
        json={
            "qr_id": "BAT",
            "shelves": {"A": 20, "B": 50, "C": 30},
            "shelf_numbers": {"A": "2", "B": "7"},
        },
    )
    pulled = client.get("/api/sync/pull").json()
    shelf = {(s["qr_id"], s["warehouse"]): s["shelf"] for s in pulled["item_stock"]}
    assert shelf[("BAT", "A")] == "2" and shelf[("BAT", "C")] == ""

    # Moving units without sending shelf numbers keeps them.
    client.post(
        "/api/panel/items/stock",
        json={"qr_id": "BAT", "shelves": {"A": 25, "B": 45, "C": 30}},
    )
    rows = client.get("/api/panel/warehouses/contents", params={"name": "A"}).json()[
        "rows"
    ]
    # Shelf 2 before shelf 10.
    assert [(r["qr_id"], r["shelf"], r["quantity"]) for r in rows] == [
        ("BAT", "2", 25),
        ("K", "10", 3),
    ]

    # The count found 23 batteries and 3 helmets in A.
    r = client.post(
        "/api/panel/warehouses/count",
        json={"name": "A", "counts": {"BAT": 23, "K": 3}},
    ).json()
    assert r["success"] and "2 פריטים, 1 עם הפרש" in r["message"]
    item, shelves = _stock(client, "BAT")
    assert shelves == {"A": 23, "B": 45, "C": 30} and item["quantity"] == 98

    assert client.get("/panel/count", params={"name": "A"}).status_code == 200
    r = client.get("/api/panel/warehouses/count.xlsx", params={"name": "A"})
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in ws[2]][:5] == ["2", "BAT", "סוללות", "סוללות", 23]
