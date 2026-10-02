import io

from openpyxl import Workbook, load_workbook


def _xlsx(sheets):
    wb = Workbook()
    del wb[wb.sheetnames[0]]
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _push(client, tx_id, qr_id, user_id, action, ts):
    r = client.post(
        "/api/sync/push",
        json={
            "device_id": "d",
            "transactions": [
                {
                    "tx_id": tx_id,
                    "qr_id": qr_id,
                    "user_id": user_id,
                    "action_type": action,
                    "timestamp": ts,
                }
            ],
        },
    )
    assert r.json()["accepted"] == [tx_id]


def test_excel_import_loads_types_soldiers_and_serials(client):
    data = _xlsx(
        {
            "סוגים": [["סוג פריט", "כמות בתקן"], ["מכשיר קשר", 3]],
            "חיילים": [["מ.א", "שם מלא", "יחידה"], [1234567, "דנה כהן", "פלוגה ב"]],
            "מלאי": [
                ["מספר סידורי", "שם פריט", "סוג פריט"],
                ["MK-1", "מכשיר קשר 710", "מכשיר קשר"],
                ["MK-2", None, "מכשיר קשר"],
                [None, "שורה ריקה", None],
            ],
        }
    )
    r = client.post("/api/panel/import", content=data, headers={"x-filename": "a.xlsx"})
    body = r.json()
    assert body["success"] is True, body
    assert body["message"] == "הייבוא הושלם: 1 חיילים, 2 פריטים, 1 סוגי פריטים"

    pulled = client.get("/api/sync/pull").json()
    items = {i["qr_id"]: i for i in pulled["items"]}
    assert items["MK-1"]["category"] == "מכשיר קשר"
    assert items["MK-2"]["name"] == "מכשיר קשר"  # name falls back to the type
    assert {"user_id": "1234567", "full_name": "דנה כהן", "unit": "פלוגה ב"} in pulled[
        "users"
    ]
    assert pulled["categories"] == [{"name": "מכשיר קשר", "target_qty": 3}]

    summary = {c["name"]: c for c in client.get("/api/panel/data").json()["categories"]}
    assert summary["מכשיר קשר"]["total"] == 2
    assert summary["מכשיר קשר"]["missing"] == 1
    assert summary[""]["label"] == "ללא סוג"  # the item from the fixture


def test_csv_import_and_bad_file(client):
    csv_data = "מספר אישי,שם מלא,יחידה\n555,יוסי לוי,מפקדה\n".encode("utf-8-sig")
    r = client.post(
        "/api/panel/import", content=csv_data, headers={"x-filename": "u.csv"}
    )
    assert r.json()["success"] is True
    r = client.post("/api/panel/import", content=b"\x00\x01garbage")
    assert r.status_code == 400
    assert r.json()["success"] is False


def test_item_and_user_management(client):
    r = client.post(
        "/api/panel/items", json={"qr_id": "S-9", "name": "", "category": "משקפת"}
    )
    assert r.json()["success"] is True
    r = client.post(
        "/api/panel/users", json={"user_id": "u9", "full_name": "רון", "unit": "ג"}
    )
    assert r.json()["success"] is True

    # Edits keep the status.
    _push(client, "t1", "S-9", "u9", "BORROW", 1_700_000_000_000)
    client.post(
        "/api/panel/items",
        json={"qr_id": "S-9", "name": "משקפת 7x50", "category": "משקפת"},
    )
    item = next(
        i for i in client.get("/api/sync/pull").json()["items"] if i["qr_id"] == "S-9"
    )
    assert item["current_status"] == "BORROWED" and item["name"] == "משקפת 7x50"

    # Items/soldiers with history cannot be deleted.
    r = client.delete("/api/panel/items", params={"qr_id": "S-9"})
    assert r.status_code == 400 and "היסטוריית" in r.json()["message"]
    r = client.delete("/api/panel/users", params={"user_id": "u9"})
    assert r.status_code == 400
    client.post("/api/panel/items", json={"qr_id": "S-10", "category": "משקפת"})
    assert client.delete("/api/panel/items", params={"qr_id": "S-10"}).json()["success"]


def test_category_rename_moves_items_and_blocks_delete(client):
    client.post("/api/panel/items", json={"qr_id": "A1", "category": "פנס"})
    client.post("/api/panel/categories", json={"name": "פנס", "target_qty": 5})
    r = client.post(
        "/api/panel/categories",
        json={"name": "פנס ראש", "target_qty": 6, "old_name": "פנס"},
    )
    assert r.json()["success"] is True
    item = next(
        i for i in client.get("/api/sync/pull").json()["items"] if i["qr_id"] == "A1"
    )
    assert item["category"] == "פנס ראש"
    r = client.delete("/api/panel/categories", params={"name": "פנס ראש"})
    assert r.status_code == 400 and "1 פריטים" in r.json()["message"]


def test_soldier_and_item_cards(client):
    _push(client, "t1", "q1", "u1", "ISSUE", 1_700_000_000_000)
    card = client.get("/api/panel/users/card", params={"user_id": "u1"}).json()
    assert card["user"]["full_name"] == "ישראל ישראלי"
    assert [h["qr_id"] for h in card["holding"]] == ["q1"]
    assert card["history"][0]["action_type"] == "ISSUE"

    _push(client, "t2", "q1", "u1", "RETURN", 1_700_000_100_000)
    item = client.get("/api/panel/items/card", params={"qr_id": "q1"}).json()
    assert item["holder"] is None
    assert [h["action_type"] for h in item["history"]] == ["RETURN", "ISSUE"]
    assert (
        client.get("/api/panel/users/card", params={"user_id": "u1"}).json()["holding"]
        == []
    )
    # Tablets get the history too.
    history = client.get("/api/sync/pull").json()["history"]
    assert {h["tx_id"] for h in history} == {"t1", "t2"}


def test_excel_exports(client):
    _push(client, "t1", "q1", "u1", "BORROW", 1_700_000_000_000)
    r = client.get("/api/panel/export/full")
    assert r.status_code == 200
    assert "filename*=UTF-8''" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == [
        "סיכום לפי סוג",
        "מלאי",
        "ציוד אצל חיילים",
        "חיילים",
        "יומן פעולות",
    ]
    assert wb["מלאי"]["D2"].value == "מושאל"
    assert wb["ציוד אצל חיילים"]["A2"].value == "ישראל ישראלי"

    r = client.get(
        "/api/panel/export/transactions", params={"start": 1_800_000_000_000}
    )
    assert load_workbook(io.BytesIO(r.content))["יומן פעולות"].max_row == 1

    r = client.get("/api/panel/users/export", params={"user_id": "u1"})
    assert load_workbook(io.BytesIO(r.content)).sheetnames == [
        "פרטי חייל",
        "ציוד אצלו",
        "היסטוריה",
    ]
    assert client.get("/api/panel/export/nope").status_code == 404
    assert client.get("/api/panel/template").status_code == 200


def test_labels_page_has_qr_codes(client):
    client.post(
        "/api/panel/items", json={"qr_id": "L-1", "name": "<b>x</b>", "category": "ק"}
    )
    page = client.get("/panel/labels", params={"ids": "L-1,q1"}).text
    assert page.count("<svg") == 2
    assert "&lt;b&gt;x&lt;/b&gt;" in page
    assert client.get("/panel/labels", params={"category": "ק"}).text.count("<svg") == 1


def test_old_database_gets_category_column(tmp_path, monkeypatch):
    import sqlite3
    import sys

    db_file = tmp_path / "old.db"
    con = sqlite3.connect(db_file)
    con.execute(
        "CREATE TABLE items (qr_id VARCHAR(128) PRIMARY KEY, name VARCHAR(200), current_status VARCHAR(32))"
    )
    con.execute("INSERT INTO items VALUES ('x', 'ישן', 'AVAILABLE')")
    con.commit()
    con.close()
    monkeypatch.setenv("WAREHOUSE_DATABASE_URL", f"sqlite:///{db_file}")
    for name in list(sys.modules):
        if name == "app" or name.startswith("app."):
            del sys.modules[name]
    from app.database import SessionLocal, init_db
    from app import services

    init_db()
    with SessionLocal() as db:
        assert services.pull_state(db).items[0].category == ""
