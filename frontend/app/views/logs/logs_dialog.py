
from PySide6.QtWidgets import (
    QDialog,
    QVBoxLayout,
    QListWidget,
    QLabel,
    QPushButton,
    QFileDialog,
)
from datetime import datetime
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from app.views.common.warning_dialog import show_warning


class LogsDialog(QDialog):
    def __init__(self, logs, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Logs")
        self.setMinimumSize(500, 400)

        if not isinstance(logs, list):
            message = logs.get("message", "No logs available")
            show_warning(
                "No Logs Found",
                "Could not retrieve logs for this visitor.\n\n"
                "Possible reasons:\n"
                "✓ Visitor has no recorded logs\n"
                "✓ Visitor was created before build bf46edd — visitor creation date may be missing\n\n"
                "Technical details:\n"
                f"• {message}",
            )
        else:
            layout = QVBoxLayout()
            self.logs_list = QListWidget()

            self.logs_data = []  # Store logs data for Excel export
            for log in logs:
                ts = datetime.fromisoformat(log["timestamp"])
                formatted = ts.strftime("%H:%M:%S %d/%m/%Y")
                entry = f"{formatted} - {log['visitor_dbid']} - {log['visitor_name']} - {log['action']}"
                self.logs_list.addItem(entry)
                self.logs_data.append(
                    {
                        "Timestamp": formatted,
                        "Visitor DBID": log["visitor_dbid"],
                        "Visitor Name": log["visitor_name"],
                        "Action": log["action"],
                    }
                )

            layout.addWidget(QLabel("Recent Logs:"))
            layout.addWidget(self.logs_list)

            # Add a button to download logs
            self.download_button = QPushButton("Download Logs as Excel")
            self.download_button.clicked.connect(self.download_logs)
            layout.addWidget(self.download_button)

            self.setLayout(layout)

    def _is_safe_path(self, path: str) -> bool:
        import os
        home_dir = os.path.expanduser("~")
        return os.path.commonpath([home_dir, os.path.abspath(path)]) == home_dir

    def download_logs(self):
        timestamp = datetime.now().strftime("logs_%Y-%m-%d_%H-%M-%S.xlsx")
        file_path, _ = QFileDialog.getSaveFileName(
            self, "Save Logs", timestamp, "Excel Files (*.xlsx)"
        )

        if file_path:
            if not self._is_safe_path(file_path):
                show_warning(
                    "Invalid File Path",
                    "The selected file path is not allowed. Please choose a location within your home directory and avoid system folders.",
                )
                return
            try:
                wb = Workbook()
                ws = wb.active
                ws.title = "Logs"

                headers = ["Timestamp", "Visitor DBID", "Visitor Name", "Action"]
                ws.append(headers)

                for log in self.logs_data:
                    ws.append([
                        log["Timestamp"],
                        log["Visitor DBID"],
                        log["Visitor Name"],
                        log["Action"]
                    ])

                for i, col in enumerate(headers, start=1):
                    ws.column_dimensions[get_column_letter(i)].width = 25

                wb.save(file_path)

            except Exception as e:
                show_warning("Error", f"Failed to save logs: {str(e)}")
